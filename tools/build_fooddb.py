"""식약처 음식 DB + 농진청 국가표준식품성분 DB(엑셀) → fooddb.json

사용법:  python tools/build_fooddb.py <음식DB.xlsx> <식품성분표.xlsx>
- 음식 DB: 식품안전나라 > 영양성분 DB 내려받기 > 음식 DB
- 원재료: 농식품올바로 > 국가표준식품성분표 > 국가표준식품성분 DB (Excel)

결과 형식: {"v": 버전, "src": 출처, "items": [[이름, 구분, 1회제공량g, kcal, 탄, 단, 지], ...]}
영양성분은 100g(ml) 기준, 값이 없으면 null.
"""
import json
import re
import sys

import openpyxl

SOURCE_PRIORITY = [  # 같은 이름의 일반 음식이 여러 출처에 있으면 앞쪽을 씀
    '외식(분석함량)', '가정식(분석 함량)', '외식(재료량 기반 산출함량)',
    '산업체급식(재료량 기반 산출 함량)', '중고등학교급식(재료량 기반 산출함량)', '초등학교급식(재료량 기반 산출 함량)',
]


def num(x):
    if x is None:
        return None
    s = str(x).strip().replace(',', '')
    if s in ('', '-', 'tr', 'Tr', 'TR'):
        return 0.0 if s.lower() == 'tr' else None
    try:
        return float(s)
    except ValueError:
        return None


def grams(x):
    m = re.search(r'([\d.]+)\s*(g|ml|mL)', str(x or ''))
    return float(m.group(1)) if m else None


def r1(x):
    return None if x is None else round(x, 1)


def load_dishes(path):
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    H = {h: i for i, h in enumerate(next(rows))}
    franchise, plain = [], {}
    for r in rows:
        name = (r[H['식품명']] or '').strip()
        kcal = num(r[H['에너지(kcal)']])
        if not name or kcal is None:
            continue
        origin = r[H['식품기원명']] or ''
        brand = (r[H['업체명']] or '').strip()
        serving = grams(r[H['식품중량']]) or grams(r[H['1인(회)분량 참고량']]) or 100
        vals = [r1(kcal), r1(num(r[H['탄수화물(g)']])), r1(num(r[H['단백질(g)']])), r1(num(r[H['지방(g)']]))]
        if '프랜차이즈' in origin and brand and brand != '해당없음':
            short = name.split('_', 1)[1] if '_' in name else name
            franchise.append([f'{brand} {short.strip()}', brand, round(serving), *vals])
        else:
            rank = SOURCE_PRIORITY.index(origin) if origin in SOURCE_PRIORITY else 99
            if name not in plain or rank < plain[name][0]:
                # "김치찌개_참치" → "김치찌개 (참치)"
                parts = [x.strip() for x in name.split('_') if x.strip()]
                label = parts[0] + (f' ({", ".join(parts[1:])})' if len(parts) > 1 else '')
                plain[name] = (rank, [label, '음식', round(serving), *vals])
    return [v[1] for v in plain.values()] + franchise


def load_raw(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = [w for w in wb.worksheets if w.title.startswith('국가표준식품성분 Database')][-1]
    out = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i < 3 or not r[3]:
            continue
        # 열: 3 식품명, 5 에너지, 7 단백질, 8 지방, 10 탄수화물
        kcal = num(r[5])
        if kcal is None:
            continue
        out.append([str(r[3]).strip(), '원재료', 100, r1(kcal), r1(num(r[10])), r1(num(r[7])), r1(num(r[8]))])
    return out, ws.title


if __name__ == '__main__':
    dishes = load_dishes(sys.argv[1])
    raw, raw_ver = load_raw(sys.argv[2])
    items = dishes + raw
    data = {
        'v': '2026-08-28 / ' + raw_ver.replace('국가표준식품성분 Database ', 'DB'),
        'src': '식품의약품안전처 식품영양성분 DB(음식), 농촌진흥청 국가표준식품성분표',
        'items': items,
    }
    with open('fooddb.json', 'w', encoding='utf-8') as fp:
        json.dump(data, fp, ensure_ascii=False, separators=(',', ':'))
    print(f'음식 {len(dishes)}개 + 원재료 {len(raw)}개 = {len(items)}개')
